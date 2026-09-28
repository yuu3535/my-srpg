# SRPGのステータス一覧（表計算ファイル）を作る。Googleスプレッドシートに取り込んだり、表をコピペしたりして使う。
#
#   py -3.12 tools/build_stat_sheet.py [出力先 .xlsx]（既定: output/SRPG_ステータス一覧.xlsx）
#
# 元データは tools/stat_sheet_data.js が集める（採用版md の表・trialStatSystem.js・abilityData.js）。
# 数字を直すときは元の文書・データを直して、作り直す（この表計算ファイルを手で直さない）。
import datetime
import json
import os
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "output", "SRPG_ステータス一覧.xlsx")

STATS = ["HP", "力", "防御", "魔攻", "魔防", "技", "速さ", "魅力"]
FONT = "Meiryo"
HEAD_FILL = PatternFill("solid", fgColor="381547")   # 黒紫の帯（UIの意匠に合わせる）
HEAD_FONT = Font(name=FONT, bold=True, color="EFD081")
SUB_FILL = PatternFill("solid", fgColor="EFE7F3")
INPUT_FILL = PatternFill("solid", fgColor="FFF2B3")
BODY = Font(name=FONT)
BLUE = Font(name=FONT, color="0000FF")   # 文書から写した数字（入力）
THIN = Side(style="thin", color="C8B6D6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "data.json")
        subprocess.run(["node", os.path.join(ROOT, "tools", "stat_sheet_data.js"), path], check=True, cwd=ROOT)
        with open(path, encoding="utf-8") as f:
            return json.load(f)


def header(ws, row, titles, widths=None):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def put(ws, row, col, value, font=BODY, fmt=None, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font, cell.border = font, BOX
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.startswith("=")):
        cell.alignment = Alignment(horizontal="right")
    return cell


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color="381547")
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, size=9, color="6B5A78")


def main():
    data = load()
    adopted, game_only, battles = data["adopted"], data["gameOnly"], data["battles"]
    wb = Workbook()
    today = datetime.date.today().isoformat()

    # ── 説明 ──
    ws = wb.active
    ws.title = "説明"
    lines = [
        ("SRPG ステータス一覧", True),
        (f"作成: {today}（tools/build_stat_sheet.py で自動作成。数字を直すときは元の文書を直して作り直す）", False),
        ("", False),
        ("シート", True),
        ("因果Lv1 … 採用版のキャラ10人の、因果Lv1の能力値・種族・所属・最初の兵種・幸運と勇気", False),
        ("成長率 … 個人成長率と、幸運・勇気の補正（(幸運＋最大勇気)÷40 の切り捨て。全能力に足す）", False),
        ("上限 … 能力上限（やりこみの到達点）", False),
        ("Lv計算 … 黄色のセルに因果Lvを入れると、その因果Lvでの能力値（成長率どおりに伸びた期待値）が出る", False),
        ("ゲームの仮キャラ … 採用版に表がなく、ゲームの中だけにある仮の値（敵・召喚・訓練のギュンター）", False),
        ("今の戦闘の値 … ゲームが今の戦闘で実際に使っている因果Lvと能力値（作成した日の値）", False),
        ("", False),
        ("注意", True),
        ("HP: ゲームは案②（TRPGのHP×1）で動いている。採用版の文書の表は案①（×2）で書かれているので、「因果Lv1」には両方を載せた。HPの案はまだ検討中", False),
        ("能力値の式: 因果Lvの値 ＝ 因果Lv1の値 ＋（個人成長率＋幸運・勇気補正）×（因果Lv−1）を四捨五入。上限を超えない（採用版 §6.1・§8）", False),
        ("兵種の補正・スキル・装備はふくまない（本人の能力値だけ）", False),
        ("最初の兵種は兵種表CSV（各キャラ兵種適正.csv の◎）から。ギュンターは表では「ベル」", False),
        ("", False),
        ("出典", True),
        ("採用版md/SRPG_CHARACTER_STAT_GROWTH_STANDARD.md（§5・§5.3 因果Lv1、§6・§6.1・§6.2 成長率、§8.2・§8.4 上限）", False),
        ("trialStatSystem.js（ゲームの仮の値・今の戦闘の値）／ abilityData.js（最初の兵種）", False),
        ("Googleスプレッドシートへ: ドライブにこのファイルを上げて「Googleスプレッドシートで開く」。表だけならシートを選んでコピペでもよい", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=text).font = Font(name=FONT, bold=bold, size=13 if r == 1 else 10, color="381547" if bold else "000000")
    ws.column_dimensions["A"].width = 120

    # ── 因果Lv1 ──
    ws = wb.create_sheet("因果Lv1")
    title(ws, "因果Lv1の能力値（採用版）", "HPはゲームで使う案②（×1）。右端に文書の案①（×2）も載せた。青い数字は採用版の文書から写したもの")
    header(ws, 4, ["キャラ", "種族", "所属", "最初の兵種"] + [f"{s}\n(案②)" if s == "HP" else s for s in STATS] + ["HP\n(案①)", "幸運", "最大勇気"],
           [14, 7, 14, 22] + [8] * 8 + [8, 7, 8])
    for i, a in enumerate(adopted):
        r = 5 + i
        put(ws, r, 1, a["name"])
        put(ws, r, 2, a["race"])
        put(ws, r, 3, a["group"])
        put(ws, r, 4, a["cls"])
        put(ws, r, 5, f"=M{r}/2")   # 案②＝案①の半分
        for k in range(1, 8):
            put(ws, r, 5 + k, a["base1"][k], BLUE)
        put(ws, r, 13, a["base1"][0], BLUE)
        put(ws, r, 14, a["luck"], BLUE)
        put(ws, r, 15, a["courage"], BLUE)
    ws["E4"].comment = Comment("案②＝案①÷2。ゲーム（テスト戦闘・訓練）はこちらで動いている", "総合担当")
    n_adopted = len(adopted)
    last = 4 + n_adopted

    # ── 成長率 ──
    ws = wb.create_sheet("成長率")
    title(ws, "個人成長率（採用版）", "1因果Lvごとに、この確率で+1（試験では期待値どおりに伸びるとして計算）。補正は全能力に足す")
    header(ws, 4, ["キャラ"] + STATS + ["合計", "幸運・勇気\n補正", "補正込み\n合計"], [14] + [8] * 8 + [9, 10, 10])
    for i, a in enumerate(adopted):
        r = 5 + i
        put(ws, r, 1, f"=因果Lv1!A{r}")
        for k in range(8):
            put(ws, r, 2 + k, a["growth"][k] / 100, BLUE, "0%")
        put(ws, r, 10, f"=SUM(B{r}:I{r})", fmt="0%")
        put(ws, r, 11, f"=INT((因果Lv1!N{r}+因果Lv1!O{r})/40)/100", fmt="0%")
        put(ws, r, 12, f"=J{r}+K{r}*8", fmt="0%")

    # ── 上限 ──
    ws = wb.create_sheet("上限")
    title(ws, "能力上限（採用版）", "HP上限はどちらのHPの案でも ×2 の式のまま（採用版 §4.3）")
    header(ws, 4, ["キャラ"] + STATS, [14] + [8] * 8)
    for i, a in enumerate(adopted):
        r = 5 + i
        put(ws, r, 1, f"=因果Lv1!A{r}")
        for k in range(8):
            put(ws, r, 2 + k, a["caps"][k], BLUE)

    # ── Lv計算 ──
    ws = wb.create_sheet("Lv計算")
    title(ws, "因果Lvでの能力値（期待値）", "黄色のセル（B3）に因果Lvを入れる。HPは案②。兵種・スキル・装備の補正はふくまない")
    ws["A3"] = "因果Lv"
    ws["A3"].font = Font(name=FONT, bold=True)
    put(ws, 3, 2, 25, BLUE, fill=INPUT_FILL)
    ws["C3"] = "← ここを変える（1〜）"
    ws["C3"].font = Font(name=FONT, size=9, color="6B5A78")
    header(ws, 5, ["キャラ"] + STATS, [14] + [8] * 8)
    ws.freeze_panes = "B6"
    for i in range(n_adopted):
        r, src = 6 + i, 5 + i
        put(ws, r, 1, f"=因果Lv1!A{src}")
        for k in range(8):
            base_col = get_column_letter(5 + k)          # 因果Lv1 の E〜L（HPは案②）
            g_col = get_column_letter(2 + k)              # 成長率 の B〜I
            cap_col = get_column_letter(2 + k)            # 上限 の B〜I
            put(ws, r, 2 + k,
                f"=ROUND(MIN(上限!{cap_col}{src},因果Lv1!{base_col}{src}+(成長率!{g_col}{src}+成長率!K{src})*($B$3-1)),0)")

    # ── ゲームの仮キャラ ──
    ws = wb.create_sheet("ゲームの仮キャラ")
    title(ws, "ゲームの中だけにある仮の値", "採用版の表にないキャラ（trialStatSystem.js）。因果Lv1の値・成長率・上限。HPは案②")
    header(ws, 4, ["キャラ", "種族"] + [f"{s}" for s in STATS] + [f"{s}\n成長率" for s in STATS] + [f"{s}\n上限" for s in STATS] + ["幸運", "最大勇気", "体格(SIZ)"],
           [30, 7] + [7] * 24 + [7, 8, 9])
    for i, g in enumerate(game_only):
        r = 5 + i
        put(ws, r, 1, g["name"])
        put(ws, r, 2, g["race"])
        for k in range(8):
            put(ws, r, 3 + k, g["base2"][k], BLUE)
            put(ws, r, 11 + k, g["growth"][k] / 100, BLUE, "0%")
            put(ws, r, 19 + k, g["caps"][k], BLUE)
        put(ws, r, 27, g["luck"], BLUE)
        put(ws, r, 28, g["courage"], BLUE)
        put(ws, r, 29, g["siz"], BLUE)
    ws.cell(row=6 + len(game_only), column=1,
            value="ギュンター（訓練）は採用版の幼ギュンターを元に訓練用に調えた値。敵3体とヒトダマは試験用の仮の値").font = Font(name=FONT, size=9, color="6B5A78")

    # ── 今の戦闘の値 ──
    ws = wb.create_sheet("今の戦闘の値")
    title(ws, "今の戦闘で使っている値", f"{today} 時点のゲームの計算そのまま（数字だけ。作り直すと更新される）。スキル・戦技は「習得に使う因果Lv」で覚えている分")
    header(ws, 4, ["戦闘", "キャラ", "因果Lv", "習得に使う\n因果Lv"] + STATS + ["最初の持ち物"], [16, 16, 8, 10] + [7] * 8 + [26])
    for i, b in enumerate(battles):
        r = 5 + i
        put(ws, r, 1, b["battle"])
        put(ws, r, 2, b["name"])
        put(ws, r, 3, b["lv"], BLUE)
        put(ws, r, 4, b["skillLv"], BLUE)
        for k in range(8):
            put(ws, r, 5 + k, b["stats"][k], BLUE)
        put(ws, r, 13, b["gear"])

    for sheet in wb.worksheets:
        sheet.sheet_view.zoomScale = 110
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
